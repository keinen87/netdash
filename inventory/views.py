import ipaddress
import platform
import subprocess

from ipaddress import IPv4Address

from django.conf import settings
from django.core.exceptions import PermissionDenied
from django.core.paginator import EmptyPage, PageNotAnInteger, Paginator
from django.db.models import Case, CharField, OuterRef, Q, Subquery, Value, When
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_POST

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from . import ldap_client
from .export_fields import EXPORT_FIELDS_BY_KEY, FIELD_GROUPS
from .host_filters import apply_host_filters
from .models import (
    Account,
    HardwareInfo,
    History,
    Host,
    IPAddress,
    Person,
    SIM,
    Software,
)


HOST_SORT_FIELDS = {
    "fio": "first_user_name",
    "kind": "kind_order",
    "domain": "domain_name",
    "ip": "first_ip_address",
    "phone": "first_user_phone",
}

HOSTS_PER_PAGE = 25


def _build_page_range(current, total, window=2):
    """Возвращает список страниц для навигатора. None = «…»."""
    if total <= 1:
        return [1]

    pages = {1, total}
    low = max(1, current - window)
    high = min(total, current + window)
    for n in range(low, high + 1):
        pages.add(n)

    result = []
    prev = None
    for n in sorted(pages):
        if prev is not None and n - prev > 1:
            result.append(None)
        result.append(n)
        prev = n
    return result


def _collect_subnets():
    """
    Возвращает отсортированный список подсетей /24 из всех IPv4-адресов.
    Подсети из settings.EXCLUDED_SUBNETS не попадают в результат.
    """
    excluded = set(getattr(settings, "EXCLUDED_SUBNETS", []))

    subnets = set()
    for ip in IPAddress.objects.values_list("address", flat=True):
        try:
            addr = IPv4Address(ip)
        except (ValueError, TypeError):
            continue
        octets = str(addr).split(".")
        subnet = f"{octets[0]}.{octets[1]}.{octets[2]}.0/24"
        if subnet in excluded:
            continue
        subnets.add(subnet)
    return sorted(subnets)


def dashboard(request):
    query = request.GET.get("q", "").strip()
    sort = request.GET.get("sort", "").strip()
    direction = request.GET.get("dir", "asc").strip()
    page_number = request.GET.get("page", "1").strip()
    kinds_selected = request.GET.getlist("kind")
    subnets_selected = request.GET.getlist("subnet")
    has_sim = request.GET.get("has_sim") == "1"

    hosts = Host.objects.all().prefetch_related("users", "softwares", "ip_addresses")

    # ← Вся фильтрация в одном месте — та же, что и в экспорте
    hosts = apply_host_filters(request, hosts)

    # Блоки «результаты поиска» по другим моделям — отдельно от фильтрации хостов
    persons = Person.objects.none()
    accounts = Account.objects.none()
    sims = SIM.objects.none()
    history = History.objects.none()
    softwares = Software.objects.none()

    if query:
        persons = Person.objects.filter(
            Q(full_name__icontains=query)
            | Q(position__icontains=query)
            | Q(phone_work__icontains=query)
            | Q(description__icontains=query)
            | Q(comment__icontains=query)
            | Q(hosts__domain_name__icontains=query)
            | Q(hosts__ip_addresses__address__icontains=query)
            | Q(hosts__ssid__icontains=query)
            | Q(hosts__serial_number__icontains=query)
            | Q(accounts__login__icontains=query)
            | Q(accounts__comment__icontains=query)
            | Q(sims__full_number__icontains=query)
            | Q(sims__icc__icontains=query)
        ).distinct()

        accounts = Account.objects.filter(
            Q(login__icontains=query)
            | Q(kind__icontains=query)
            | Q(comment__icontains=query)
            | Q(person__full_name__icontains=query)
            | Q(person__position__icontains=query)
            | Q(host__domain_name__icontains=query)
            | Q(host__ip_addresses__address__icontains=query)
        ).distinct()

        sims = SIM.objects.filter(
            Q(full_number__icontains=query)
            | Q(short_number__icontains=query)
            | Q(icc__icontains=query)
            | Q(apn__icontains=query)
            | Q(comment__icontains=query)
            | Q(person__full_name__icontains=query)
            | Q(host__domain_name__icontains=query)
            | Q(host__ip_addresses__address__icontains=query)
        ).distinct()

        history = History.objects.filter(
            Q(fact__icontains=query)
            | Q(host__domain_name__icontains=query)
            | Q(host__ip_addresses__address__icontains=query)
            | Q(person__full_name__icontains=query)
        ).distinct()

        softwares = Software.objects.filter(
            Q(name__icontains=query)
            | Q(version__icontains=query)
            | Q(comment__icontains=query)
            | Q(hosts__domain_name__icontains=query)
            | Q(hosts__ip_addresses__address__icontains=query)
        ).distinct()

    # Виртуальное поле для сортировки по русскому названию типа
    kind_order_case = Case(
        *[
            When(kind=code, then=Value(label.lower()))
            for code, label in Host.KIND_CHOICES
        ],
        output_field=CharField(),
        default=Value(""),
    )

    first_user = (
        Person.objects
        .filter(hosts=OuterRef("pk"))
        .order_by("full_name")
    )
    first_ip = (
        IPAddress.objects
        .filter(host=OuterRef("pk"))
        .order_by("-is_primary", "address")
    )
    hosts = hosts.annotate(
        first_user_name=Subquery(first_user.values("full_name")[:1]),
        first_user_phone=Subquery(first_user.values("phone_work")[:1]),
        first_ip_address=Subquery(first_ip.values("address")[:1]),
        kind_order=kind_order_case,
    )

    if sort in HOST_SORT_FIELDS:
        field = HOST_SORT_FIELDS[sort]
        if direction == "desc":
            field = "-" + field
        hosts = hosts.order_by(field)

    paginator = Paginator(hosts, HOSTS_PER_PAGE)
    try:
        page_obj = paginator.page(page_number)
    except PageNotAnInteger:
        page_obj = paginator.page(1)
    except EmptyPage:
        page_obj = paginator.page(paginator.num_pages)

    query_params = request.GET.copy()
    query_params.pop("page", None)
    query_suffix = query_params.urlencode()
    if query_suffix:
        query_suffix += "&"

    sort_params = request.GET.copy()
    sort_params.pop("page", None)
    sort_params.pop("sort", None)
    sort_params.pop("dir", None)
    sort_suffix = sort_params.urlencode()
    if sort_suffix:
        sort_suffix += "&"

    context = {
        "query": query,
        "sort": sort,
        "direction": direction,
        "kinds_selected": kinds_selected,
        "kind_choices": Host.KIND_CHOICES,
        "subnets_selected": subnets_selected,
        "available_subnets": _collect_subnets(),
        "has_sim": has_sim,
        "page_obj": page_obj,
        "paginator": paginator,
        "page_range": _build_page_range(page_obj.number, paginator.num_pages),
        "query_suffix": query_suffix,
        "sort_suffix": sort_suffix,
        "hosts": page_obj,
        "persons": persons,
        "accounts": accounts,
        "sims": sims,
        "history": history,
        "softwares": softwares,
    }
    return render(request, "inventory/dashboard.html", context)


def host_detail(request, pk):
    host = get_object_or_404(
        Host.objects.prefetch_related("ip_addresses", "users", "softwares"),
        pk=pk,
    )

    can_edit = request.user.is_authenticated and request.user.is_staff

    if request.method == "POST":
        if not can_edit:
            raise PermissionDenied("Добавлять записи в историю может только администратор.")

        fact = request.POST.get("fact", "").strip()
        if fact:
            default_person = host.users.first()
            History.objects.create(
                host=host,
                person=default_person,
                fact=fact,
            )
        return redirect("inventory:host_detail", pk=host.pk)

    hardware = HardwareInfo.objects.filter(host=host).first()

    context = {
        "host": host,
        "hardware": hardware,
        "users": host.users.all(),
        "ip_addresses": host.ip_addresses.all(),
        "accounts": host.accounts.select_related("person").all(),
        "sims": host.sims.select_related("person").all(),
        "softwares": host.softwares.all(),
        "history": host.history.all(),
        "can_edit": can_edit,
        "web_ui_url": host.web_ui_url,
    }
    return render(request, "inventory/host_detail.html", context)


@require_POST
def host_ping(request, pk):
    if not (request.user.is_authenticated and request.user.is_staff):
        return JsonResponse(
            {"ok": False, "error": "Проверка связи доступна только администратору."},
            status=403,
        )

    host = get_object_or_404(Host, pk=pk)
    ip = host.primary_ip

    if not ip:
        return JsonResponse(
            {"ok": False, "error": "У хоста не задан IP-адрес."},
            status=400,
        )

    try:
        ipaddress.ip_address(ip)
    except ValueError:
        return JsonResponse({"ok": False, "error": "Некорректный IP-адрес."}, status=400)

    system = platform.system().lower()
    if system == "windows":
        cmd = ["ping", "-n", "4", "-w", "2000", ip]
    else:
        cmd = ["ping", "-c", "4", "-W", "2", ip]

    try:
        result = subprocess.run(cmd, capture_output=True, timeout=15)
    except subprocess.TimeoutExpired:
        return JsonResponse({"ok": False, "error": "Тайм-аут ping (15 секунд)."})
    except FileNotFoundError:
        return JsonResponse({"ok": False, "error": "Утилита ping не найдена на сервере."})
    except Exception as exc:
        return JsonResponse({"ok": False, "error": f"Ошибка: {exc}"})

    raw = (result.stdout or b"") + (result.stderr or b"")
    if system == "windows":
        candidates = ("cp866", "cp1251", "utf-8")
    else:
        candidates = ("utf-8",)

    output = None
    for enc in candidates:
        try:
            output = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if output is None:
        output = raw.decode("utf-8", errors="replace")

    return JsonResponse({"ok": result.returncode == 0, "output": output.strip()})


def ldap_search_view(request):
    """Поиск в LDAP-каталоге. Доступ — только администраторам."""
    if not (request.user.is_authenticated and request.user.is_staff):
        raise PermissionDenied("Поиск в каталоге доступен только администратору.")

    query = request.GET.get("q", "").strip()
    results = []
    error = None

    if query:
        results, error = ldap_client.search_ldap(query)

    context = {
        "query": query,
        "results": results,
        "error": error,
        "is_configured": ldap_client.is_configured(),
        "attribute_labels": ldap_client.ATTRIBUTE_LABELS,
        "attribute_order": [
            "domain_login",
            "displayName",
            "givenName",
            "sn",
            "mail",
            "telephoneNumber",
            "mobile",
            "title",
            "department",
            "company",
            "physicalDeliveryOfficeName",
            "whenCreated",
            "lastLogonTimestamp",
        ],
    }
    return render(request, "inventory/ldap_search.html", context)


def export_page(request):
    """Страница выбора полей выгрузки. Текущие фильтры пробрасываются в форму."""
    return render(request, "inventory/export.html", {
        "field_groups": FIELD_GROUPS,
        "export_fields_by_key": EXPORT_FIELDS_BY_KEY,
        "query": request.GET.get("q", ""),
        "kinds_selected": request.GET.getlist("kind"),
        "subnets_selected": request.GET.getlist("subnet"),
        "has_sim": request.GET.get("has_sim") == "1",
    })


def export_xlsx(request):
    """Генерирует XLSX по выбранным полям и текущим фильтрам."""
    selected = request.GET.getlist("field")
    # whitelist: только известные ключи
    selected = [k for k in selected if k in EXPORT_FIELDS_BY_KEY]
    if not selected:
        selected = ["kind", "domain_name", "primary_ip"]  # дефолт

    hosts = Host.objects.prefetch_related(
        "users", "ip_addresses", "softwares", "sims"
    )
    hosts = apply_host_filters(request, hosts)

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Хосты")

    # Закрепляем первую строку — заголовки всегда видны
    ws.freeze_panes = "A2"

    # Ширина столбцов — по заголовкам, с запасом
    for idx, key in enumerate(selected, start=1):
        label = EXPORT_FIELDS_BY_KEY[key][0]
        ws.column_dimensions[get_column_letter(idx)].width = max(14, len(label) + 4)

    # Заголовки жирным
    bold = Font(bold=True)
    header_cells = []
    for key in selected:
        label = EXPORT_FIELDS_BY_KEY[key][0]
        cell = WriteOnlyCell(ws, value=label)
        cell.font = bold
        header_cells.append(cell)
    ws.append(header_cells)

    # Данные
    for host in hosts.iterator(chunk_size=500):
        row = []
        for key in selected:
            label, getter = EXPORT_FIELDS_BY_KEY[key]
            try:
                value = getter(host)
            except Exception:
                value = ""
            row.append(value)
        ws.append(row)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    filename = f"hosts_export_{timezone.now():%Y%m%d_%H%M%S}.xlsx"
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response